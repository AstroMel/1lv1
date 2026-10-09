"""Writes the ranked list as a new Excel file. Never touches the input workbook."""
import hashlib

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import scoring

COLUMNS = [
    ("Rank", "rank", 7, "0"), ("Visit list", "visit_list", 11, None),
    ("Dealer ID", "dealer_id", 11, None), ("Dealer", "dealer_name", 26, None),
    ("Region", "region", 14, None), ("Priority (0-100)", "priority", 11, "0.0"),
    ("Why", "why", 70, None),
    ("24-mo spend", "spend_24m", 13, '"$"#,##0'), ("Last 90d spend", "recent_90d", 13, '"$"#,##0'),
    ("Usual 90d spend", "usual_90d", 13, '"$"#,##0'),
    ("Last purchase", "last_purchase", 13, "yyyy-mm-dd"),
    ("Days since purchase", "days_since_purchase", 11, "0"),
    ("Buy rhythm (days)", "buy_rhythm_days", 11, "0"),
    ("Last visit", "last_visit", 13, "yyyy-mm-dd"), ("Days since visit", "days_since_visit", 11, "0"),
    ("Value score", "value", 9, "0"), ("Decline score", "decline", 9, "0"),
    ("Lapse score", "lapse", 9, "0"), ("Visit-gap score", "visit_gap", 9, "0"),
]
FILL = {"Visit now": "F8CBAD", "Next": "FFE699", "Maintain": "E2EFDA"}


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _val(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    return v.to_pydatetime() if isinstance(v, pd.Timestamp) else v


def write(path, ranked, issues, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Priority"
    head = PatternFill("solid", fgColor="1F3864")
    for c, (label, _, width, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, rec in enumerate(ranked.to_dict("records"), 2):
        for c, (_, key, _, fmt) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=r, column=c, value=_val(rec[key]))
            if fmt:
                cell.number_format = fmt
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=FILL[rec["visit_list"]])
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32

    wi = wb.create_sheet("Data checks")
    wi.append(["Sheet", "Row", "Dealer ID", "Problem"])
    for i in issues:
        wi.append(list(i))
    if not issues:
        wi.append(["", "", "", "No problems found in the input data"])
    for col, wd in zip("ABCD", (12, 7, 12, 90)):
        wi.column_dimensions[col].width = wd
    wi["A1"].font = wi["B1"].font = wi["C1"].font = wi["D1"].font = Font(bold=True)

    wm = wb.create_sheet("Method & authorization")
    w = meta["scoring"]["weights"]
    lines = [
        ("DEALER VISIT PRIORITY: how this list was made", True),
        (f"As-of date: {meta['as_of']}    Input: {meta['input_name']}    "
         f"Input SHA-256: {meta['input_sha256']}", False),
        (f"Dealers ranked: {len(ranked)}    Data problems flagged: {len(issues)} "
         "(see 'Data checks'; those rows were not used)", False),
        ("", False),
        ("AUTHORIZATION", True),
        (f"Agent owner: {meta['owner']}    Owner key fingerprint: {meta['fingerprint']}", False),
        ("Permissions used in this run: " + "; ".join(
            f"{a} on {t}" for a, t in meta["used"]), False),
        ("This tool did not send, receive, or change anything. Sending, receiving and editing "
         "source files are not built in and cannot be switched on without a new, owner-signed "
         "version of the tool.", False),
        ("The ranking rules below are signed by the owner. No individual can adjust them.", False),
        ("", False),
        ("HOW THE SCORE WORKS (each part is 0-100; higher = visit sooner)", True),
        (f"Value ({w['value']:.0%}): where the dealer ranks among all dealers by auction spend over "
         f"the last {scoring.VALUE_DAYS} days.", False),
        (f"Decline ({w['decline']:.0%}): how far spend in the last {scoring.RECENT_DAYS} days is "
         "below the dealer's own usual 90-day pace (the 365 days before that). Growth scores 0.", False),
        (f"Lapse ({w['lapse']:.0%}): days since last purchase compared to the dealer's own median "
         f"gap between purchases (default {meta['scoring']['default_buy_interval_days']} days "
         "if fewer than 3 purchase days). 0 until overdue; 100 at 3x their rhythm.", False),
        (f"Visit gap ({w['visit_gap']:.0%}): days since last visit vs a {meta['scoring']['visit_interval_days']}"
         f"-day target. 100 at 2x the target or if no visit is on record.", False),
        ("Priority = weighted sum of the four scores. Ties break by 24-month spend, then dealer ID, "
         "so the order is always the same for the same inputs.", False),
        (f"'Visit now' = rank 1-{meta['scoring']['top_n']}; 'Next' = the following "
         f"{meta['scoring']['top_n']}; 'Maintain' = everyone else.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = wm.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wm.column_dimensions["A"].width = 140
    wb.save(path)
