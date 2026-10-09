import json
import os
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import make_sample
import permissions
import run
import scoring
from permissions import Gate, PermissionDenied

AS_OF = "2026-10-09"


def owner_home(tmp, grants=()):
    home = Path(tmp) / "home"
    key = Path(tmp) / "owner.key"
    permissions.main(["--home", str(home), "keygen", "--key-out", str(key), "--owner", "Test Owner"])
    for action, scope in grants:
        permissions.main(["--home", str(home), "grant", "--key", str(key), "--action", action,
                          "--scope", scope, "--days", "30"])
    return home, key


def write_sample(tmp, n=40):
    d, p = make_sample.build(n)
    path = Path(tmp) / "dealers.xlsx"
    with pd.ExcelWriter(path) as w:
        d.to_excel(w, sheet_name="Dealers", index=False)
        p.to_excel(w, sheet_name="Purchases", index=False)
    return path


class PermissionTests(unittest.TestCase):
    def setUp(self):
        self._t = tempfile.TemporaryDirectory()
        self.tmp = self._t.name

    def tearDown(self):
        self._t.cleanup()

    def test_default_denies_everything(self):
        g = Gate(Path(self.tmp) / "nothing")
        with self.assertRaises(PermissionDenied):
            g.require("read_workbook", "dealers.xlsx")

    def test_fresh_owner_policy_has_zero_grants(self):
        home, _ = owner_home(self.tmp)
        with self.assertRaises(PermissionDenied):
            Gate(home).require("read_workbook", "dealers.xlsx")

    def test_grant_allows_only_matching_file_and_action(self):
        home, _ = owner_home(self.tmp, [("read_workbook", "dealers.xlsx")])
        g = Gate(home)
        g.require("read_workbook", "/any/dir/dealers.xlsx")
        with self.assertRaises(PermissionDenied):
            g.require("read_workbook", "other.xlsx")
        with self.assertRaises(PermissionDenied):
            g.require("write_report", "dealers.xlsx")

    def test_send_receive_edit_cannot_be_granted_or_used(self):
        home, key = owner_home(self.tmp)
        for a in ("send_message", "receive_message", "edit_source_data"):
            with self.assertRaises(SystemExit):
                permissions.main(["--home", str(home), "grant", "--key", str(key),
                                  "--action", a, "--scope", "*", "--days", "1"])
            with self.assertRaises(PermissionDenied):
                Gate(home).require(a, "x")

    def test_expired_grant_denied(self):
        home, _ = owner_home(self.tmp, [("read_workbook", "*.xlsx")])
        later = datetime(2030, 1, 1, tzinfo=timezone.utc)
        with self.assertRaises(PermissionDenied):
            Gate(home, now=later).require("read_workbook", "a.xlsx")

    def test_operator_cannot_edit_policy(self):
        home, _ = owner_home(self.tmp)
        pol = json.loads((home / "policy.json").read_text())
        pol["policy"]["grants"].append({"action": "read_workbook", "scope": "*",
                                        "expires": "2099-01-01", "note": "self-granted"})
        (home / "policy.json").write_text(json.dumps(pol))
        g = Gate(home)
        self.assertIsNotNone(g.problem)
        with self.assertRaises(PermissionDenied):
            g.require("read_workbook", "a.xlsx")

    def test_someone_elses_key_cannot_sign(self):
        home, _ = owner_home(self.tmp)
        other = Path(self.tmp) / "other.key"
        other_home = Path(self.tmp) / "other_home"
        permissions.main(["--home", str(other_home), "keygen", "--key-out", str(other),
                          "--owner", "Imposter"])
        with self.assertRaises(SystemExit):
            permissions.main(["--home", str(home), "grant", "--key", str(other), "--action",
                              "read_workbook", "--scope", "*", "--days", "1"])

    def test_scoring_rules_change_only_via_owner_and_must_sum_to_one(self):
        home, key = owner_home(self.tmp)
        with self.assertRaises(SystemExit):
            permissions.main(["--home", str(home), "set-scoring", "--key", str(key),
                              "--weights", "value=.9,decline=.9,lapse=.1,visit_gap=.1"])

    def test_audit_log_is_tamper_evident(self):
        home, _ = owner_home(self.tmp, [("read_workbook", "a.xlsx")])
        g = Gate(home)
        g.require("read_workbook", "a.xlsx")
        with self.assertRaises(PermissionDenied):
            g.require("write_report", "a.xlsx")
        self.assertEqual(permissions.verify_audit(home), (True, 2))
        log = home / "audit.log.jsonl"
        log.write_text(log.read_text().replace('"deny"', '"allow"'))
        self.assertFalse(permissions.verify_audit(home)[0])


class RunTests(unittest.TestCase):
    def setUp(self):
        self._t = tempfile.TemporaryDirectory()
        self.tmp = self._t.name
        self.wb = write_sample(self.tmp)

    def tearDown(self):
        self._t.cleanup()

    def test_run_without_permission_creates_nothing(self):
        home, _ = owner_home(self.tmp)
        out = Path(self.tmp) / "out"
        with self.assertRaises(PermissionDenied):
            run.run(self.wb, out, AS_OF, home)
        self.assertFalse(out.exists())

    def test_read_without_write_permission_creates_nothing(self):
        home, _ = owner_home(self.tmp, [("read_workbook", "dealers.xlsx")])
        out = Path(self.tmp) / "out"
        with self.assertRaises(PermissionDenied):
            run.run(self.wb, out, AS_OF, home)
        self.assertFalse(out.exists())

    def test_full_run_is_deterministic_and_order_independent(self):
        home, _ = owner_home(self.tmp, [("read_workbook", "*.xlsx"),
                                        ("write_report", "priority_*.xlsx")])
        path, ranked, issues = run.run(self.wb, Path(self.tmp) / "out", AS_OF, home)
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Priority", "Data checks", "Method & authorization"])
        self.assertEqual(len(issues), 3)
        self.assertIn("did not send, receive, or change anything",
                      " ".join(str(c.value) for r in wb["Method & authorization"].iter_rows()
                               for c in r if c.value))
        # same data, shuffled rows -> identical ranking
        sheets = pd.read_excel(self.wb, sheet_name=None)
        shuffled = {k: v.sample(frac=1, random_state=3).reset_index(drop=True)
                    for k, v in sheets.items()}
        cfg = Gate(home).policy["scoring"]
        a, _ = scoring.compute(sheets, AS_OF, cfg)
        b, _ = scoring.compute(shuffled, AS_OF, cfg)
        self.assertEqual(list(a.dealer_id), list(b.dealer_id))
        self.assertEqual(list(a.priority), list(b.priority))

    def test_source_workbook_untouched_and_never_overwritten(self):
        home, _ = owner_home(self.tmp, [("read_workbook", "*.xlsx"),
                                        ("write_report", "*.xlsx")])
        before = Path(self.wb).read_bytes()
        run.run(self.wb, Path(self.tmp) / "out", AS_OF, home)
        self.assertEqual(before, Path(self.wb).read_bytes())


class ScoringTests(unittest.TestCase):
    CFG = permissions.DEFAULT_SCORING

    def build(self):
        dealers = pd.DataFrame({
            "Dealer ID": [1, 2, 3, 4], "Name": ["Big dropping", "Big steady", "Small", "Never bought"],
            "Region": ["N"] * 4,
            "Last Visit": pd.to_datetime(["2026-04-01", "2026-09-20", "2026-09-20", None])})
        rows = []
        for d in pd.date_range("2025-10-15", "2026-06-15", freq="14D"):
            rows.append((1, d, 20000))
        for d in pd.date_range("2025-10-15", "2026-10-01", freq="14D"):
            rows.append((2, d, 20000))
            rows.append((3, d, 500))
        p = pd.DataFrame(rows, columns=["dealer_id", "Purchase Date", "Amount"])
        return {"Dealers": dealers, "Purchases": p}

    def test_ordering_makes_business_sense(self):
        r, issues = scoring.compute(self.build(), AS_OF, self.CFG)
        self.assertEqual(issues, [])
        self.assertEqual(r.dealer_id.iloc[0], "1")           # big, stopped buying, not visited
        pos = {d: i for i, d in enumerate(r.dealer_id)}
        self.assertLess(pos["2"], pos["3"])                  # same rhythm, bigger spend first
        self.assertEqual(r.set_index("dealer_id").loc["4", "why"].startswith("No purchases"), True)

    def test_missing_column_fails_loudly(self):
        s = self.build()
        s["Purchases"] = s["Purchases"].drop(columns=["Amount"])
        with self.assertRaises(scoring.InputError):
            scoring.compute(s, AS_OF, self.CFG)


if __name__ == "__main__":
    unittest.main()
